import io

from openpyxl import load_workbook

from app.parsers.blocks import Block


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_xlsx(data: bytes) -> list[Block]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks = []
    for ws in wb.worksheets:
        rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if len(rows) < 2:
            continue
        header = rows[0]
        for n, row in enumerate(rows[1:], start=2):
            text = "; ".join(f"{h}: {v}" for h, v in zip(header, row) if h and v)
            blocks.append(Block("row", f"{ws.title}. {text}", section=ws.title, sheet=ws.title, row=n))
    wb.close()
    return blocks
